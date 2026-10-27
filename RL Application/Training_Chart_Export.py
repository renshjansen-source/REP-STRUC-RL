"""
Export TensorBoard training metrics as square PNG charts, driven by Training_Data_Overview.xlsx.

Excel layout (header row is found by the cell "Title"):
- "Title"          -> chart title (single runs) / legend entry (grouped runs)
- "ID"             -> run folder relative to LOG_ROOT, e.g. "20260805_085217/PPO_1"
- "Grouping"       -> optional label; runs sharing a label are plotted together, per metric
- "Desired Log *"  -> TensorBoard scalar tags to plot (duplicates within a row are ignored)

Output:
- Single run : OUTPUT_ROOT/<Title>/<Title>_<metric>.png        + <ID_FILE> with the run ID
- Group      : OUTPUT_ROOT/Group <label>/Group <label>_<metric>.png + <ID_FILE> with "Title: ID" per run
- Flat copy  : OUTPUT_ROOT/<FLAT_FOLDER>/<same file name>.png (all images, no subfolders)
- "/" in tags becomes "_"; characters illegal in Windows file names are removed.
- Folder names must be unique (case-insensitive, as on Windows) -> duplicate titles are reported as issues.

Before anything is saved, every run folder, its event files, its requested tags and the output path
lengths are checked. If anything is missing, the issues are printed and NO images are saved.
"""

import os
import sys
import glob
import re
import textwrap

import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from tensorboard.backend.event_processing.event_accumulator import EventAccumulator, SCALARS

# =============================================================================
# CONFIG
# =============================================================================

EXCEL_PATH  = r"C:\Users\rensh\Documents\Master Thesis\Report\Images and Charts\Appendices\Appendix C - Training Data\Training Data Overview 2.xlsx"
LOG_ROOT    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logs")
OUTPUT_ROOT = r"C:\Users\rensh\Documents\Master Thesis\Report\Images and Charts\Appendices\Appendix C - Training Data"
FLAT_FOLDER = "_All Charts"   # every image again, without subfolders
ID_FILE     = "run_id.txt"    # written into each run / group folder

FIG_SIZE_IN    = 6.0    # square figure, inches
DPI            = 300
TITLE_WRAP     = 48     # characters per title line
LEGEND_COLUMNS = 2
MAX_PATH_LEN   = 259    # Windows MAX_PATH (260) minus terminating null
CHECK_PATH_LEN = True   # set False only if Windows long paths are enabled (LongPathsEnabled registry key)

X_LABEL = "Timesteps"

# Y-axis labels per tag. Units follow BikeBuilder_Callback.py / BikeBuilder_FEA.py.
METRIC_LABELS = {
    "penalties/intersection_ep_percentage" : "Intersecting steps per episode (%)",
    "penalties/reuse_ep_percentage"        : "Frame reuse steps per episode (%)",
    "termination/termination_ep_percentage": "Terminated episodes per rollout (%)",
    "termination/true_termination_tally"   : "True terminations, cumulative (count)",
    "termination/true_termination_ep_percentage": "True terminated episodes per rollout (%)",
    "train/explained_variance"             : "Explained variance (-)",
    "time/fps"                             : "Throughput (timesteps/s)",
    "behaviour/placed_frames_mean"         : "Placed frames per episode (mean)",
    "rollout/ep_len_mean"                  : "Episode length (steps, mean)",
    "FEA/max_displacement_mean"            : "Max. displacement (cm, mean)",
}

# Rows listing BOTH trigger tags (runs logged before true_termination_ep_percentage existed) get one
# reconstructed DERIVED_TAG chart instead of the two trigger charts. Computed from DERIVE_SOURCES.
DERIVE_TRIGGER = ("termination/true_termination_tally", "rollout/ep_len_mean")
DERIVE_SOURCES = ("termination/true_termination_tally", "termination/episode_tally")
DERIVED_TAG    = "termination/true_termination_ep_percentage"

# Categorical order, fixed (dataviz reference palette, light mode). 8 slots = max group size.
SERIES_COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]

# =============================================================================
# EXCEL
# =============================================================================

def read_overview(excel_path):
    """Returns list of runs: {"title", "id", "group", "metrics"} in sheet order."""
    ws = openpyxl.load_workbook(excel_path, data_only=True).active

    header_row, columns = None, {}
    for row in ws.iter_rows():
        if any(cell.value == "Title" for cell in row):
            header_row = row[0].row
            for cell in row:
                if isinstance(cell.value, str):
                    columns.setdefault(cell.value.strip(), []).append(cell.column - 1)
            break
    if header_row is None:
        sys.exit(f"No header cell 'Title' found in '{excel_path}'.")

    for required in ("Title", "ID", "Grouping"):
        if required not in columns:
            sys.exit(f"Column '{required}' not found in '{excel_path}'.")
    title_col, id_col, group_col = columns["Title"][0], columns["ID"][0], columns["Grouping"][0]
    metric_cols = sorted(c for header, cols in columns.items() if header.startswith("Desired Log") for c in cols)

    runs = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        run_id = row[id_col]
        if run_id is None or str(run_id).strip() == "":
            continue
        metrics = []
        for c in metric_cols:
            tag = row[c]
            if tag is not None and str(tag).strip() and str(tag).strip() not in metrics:
                metrics.append(str(tag).strip())
        # Rows listing both DERIVE_TRIGGER tags get DERIVED_TAG instead of those two charts
        if all(tag in metrics for tag in DERIVE_TRIGGER):
            metrics[metrics.index(DERIVE_TRIGGER[0])] = DERIVED_TAG
            metrics.remove(DERIVE_TRIGGER[1])
        group = row[group_col]
        runs.append({
            "title"  : str(row[title_col]).strip() if row[title_col] is not None else "",
            "id"     : str(run_id).strip(),
            "group"  : str(group).strip() if group is not None and str(group).strip() else None,
            "metrics": metrics,
        })
    return runs

# =============================================================================
# TENSORBOARD
# =============================================================================

def find_event_files(run_dir):
    return sorted(glob.glob(os.path.join(run_dir, "**", "events.out.tfevents.*"), recursive=True))


def derive_true_termination_percentage(scalars):
    """
    Per-rollout true termination percentage from the cumulative tallies, matching
    BikeBuilder_Callback: 100 * true_terminated_episodes / finished_episodes (both reset each rollout).
    Rollout values = differences between consecutive logged tallies.
    """
    true_tally, episode_tally = (scalars[tag] for tag in DERIVE_SOURCES)
    series, prev_true, prev_eps = {}, 0.0, 0.0
    for step in sorted(set(true_tally) & set(episode_tally)):
        if episode_tally[step] < prev_eps:            # tallies restarted (new callback) -> count from 0
            prev_true, prev_eps = 0.0, 0.0
        d_true, d_eps = true_tally[step] - prev_true, episode_tally[step] - prev_eps
        if d_eps > 0:                                  # no finished episodes -> callback logs nothing either
            series[step] = 100.0 * d_true / d_eps
        prev_true, prev_eps = true_tally[step], episode_tally[step]
    return series


def load_scalars(event_files):
    """Returns {tag: {step: value}} merged over all event files of a run."""
    scalars = {}
    for f in event_files:
        acc = EventAccumulator(f, size_guidance={SCALARS: 0})  # 0 = load all points
        acc.Reload()
        for tag in acc.Tags()["scalars"]:
            series = scalars.setdefault(tag, {})
            for event in acc.Scalars(tag):
                series[event.step] = event.value  # later files overwrite duplicate steps
    return scalars

# =============================================================================
# NAMING
# =============================================================================

def clean_name(text):
    text = text.replace("/", "_").replace("\\", "_")
    text = re.sub(r'[<>:"|?*]', "", text)
    return re.sub(r"\s+", " ", text).strip(" .")


def build_jobs(runs):
    """One job per output folder: a single ungrouped run, or all runs sharing a group label."""
    jobs, groups = [], {}
    for run in runs:
        if run["group"] is None:
            jobs.append({"name": clean_name(run["title"]), "runs": [run], "grouped": False})
        elif run["group"] not in groups:
            groups[run["group"]] = {"name": clean_name(f"Group {run['group']}"), "runs": [run], "grouped": True}
            jobs.append(groups[run["group"]])
        else:
            groups[run["group"]]["runs"].append(run)

    for job in jobs:
        job["metrics"] = []
        for run in job["runs"]:
            for tag in run["metrics"]:
                if tag not in job["metrics"]:
                    job["metrics"].append(tag)
    return jobs


def image_path(job, tag):
    return os.path.join(OUTPUT_ROOT, job["name"], f"{job['name']}_{clean_name(tag)}.png")


def flat_image_path(job, tag):
    return os.path.join(OUTPUT_ROOT, FLAT_FOLDER, os.path.basename(image_path(job, tag)))


def id_file_path(job):
    return os.path.join(OUTPUT_ROOT, job["name"], ID_FILE)

# =============================================================================
# CHECK
# =============================================================================

def check_all(runs, jobs):
    """Loads every run. Returns (scalars per run ID, list of issue strings)."""
    data, issues = {}, []

    for run in runs:
        label = f"{run['title'] or '<no title>'} [{run['id']}]"
        if not run["title"]:
            issues.append(f"{label}: no title in column 'Title'")
        if not run["metrics"]:
            issues.append(f"{label}: no desired logs listed")

        run_dir = os.path.join(LOG_ROOT, *run["id"].split("/"))
        if not os.path.isdir(run_dir):
            issues.append(f"{label}: log folder missing -> {run_dir}")
            continue
        event_files = find_event_files(run_dir)
        if not event_files:
            issues.append(f"{label}: no event files in -> {run_dir}")
            continue

        scalars = load_scalars(event_files)
        missing = False
        for tag in run["metrics"]:
            for source in (DERIVE_SOURCES if tag == DERIVED_TAG else (tag,)):
                if source not in scalars or not scalars[source]:
                    issues.append(f"{label}: metric missing -> {source}" + (f" (needed for {DERIVED_TAG})" if tag == DERIVED_TAG else ""))
                    missing = True
        if DERIVED_TAG in run["metrics"] and not missing:
            scalars[DERIVED_TAG] = derive_true_termination_percentage(scalars)  # replaces a logged one, if present
        data[run["id"]] = scalars

    # The same ID may appear as a single run AND in a group, but not twice in one group
    for job in jobs:
        ids = [r["id"] for r in job["runs"]]
        for run_id in sorted({i for i in ids if ids.count(i) > 1}):
            issues.append(f"{job['name']}: ID listed {ids.count(run_id)}x in the same group -> {run_id}")

    folders = {}
    for job in jobs:
        folders.setdefault(job["name"].lower(), []).append(job)
    for same_name in folders.values():
        if len(same_name) > 1:
            ids = ", ".join(r["id"] for job in same_name for r in job["runs"])
            issues.append(f"{same_name[0]['name']}: folder name used {len(same_name)}x -> rename titles of {ids}")
    if FLAT_FOLDER.lower() in folders:
        issues.append(f"{FLAT_FOLDER}: a title or group has the same name as FLAT_FOLDER")

    for job in jobs:
        if not job["name"]:
            issues.append(f"{job['runs'][0]['id']}: title is empty after removing illegal characters")
        if job["grouped"] and len(job["runs"]) > len(SERIES_COLORS):
            issues.append(f"{job['name']}: {len(job['runs'])} runs, max {len(SERIES_COLORS)} per group")
        for tag in job["metrics"]:
            for path in (image_path(job, tag), flat_image_path(job, tag)):
                if CHECK_PATH_LEN and len(path) > MAX_PATH_LEN:
                    issues.append(f"{job['name']}: output path {len(path)} chars (> {MAX_PATH_LEN}) -> {path}")

    return data, issues

# =============================================================================
# PLOT
# =============================================================================

def format_timesteps(x, _pos):
    """200000 -> 200k, 2000000 -> 2m, 1500000 -> 1.5m"""
    if abs(x) >= 1e6:
        return f"{x / 1e6:g}m"
    if abs(x) >= 1e3:
        return f"{x / 1e3:g}k"
    return f"{x:g}"


def new_square_figure():
    fig, ax = plt.subplots(figsize=(FIG_SIZE_IN, FIG_SIZE_IN), layout="constrained")
    ax.grid(True, color="#d9d9d9", linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.set_xlabel(X_LABEL)
    ax.xaxis.set_major_formatter(FuncFormatter(format_timesteps))
    return fig, ax


def plot_series(ax, series, color, label=None):
    steps = sorted(series)
    ax.plot(steps, [series[s] for s in steps], color=color, linewidth=1.5, label=label)


def save_figure(fig, paths):
    for path in paths:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        fig.savefig(path, dpi=DPI, facecolor="white")
        print(f"Saved: {path}")
    plt.close(fig)


def write_id_file(job):
    path = id_file_path(job)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        if job["grouped"]:
            f.writelines(f"{run['title']}: {run['id']}\n" for run in job["runs"])
        else:
            f.write(f"{job['runs'][0]['id']}\n")
    print(f"Saved: {path}")


def export_job(job, data):
    write_id_file(job)
    for tag in job["metrics"]:
        fig, ax = new_square_figure()
        ax.set_ylabel(METRIC_LABELS.get(tag, tag))

        if job["grouped"]:
            for i, run in enumerate(job["runs"]):
                if tag in run["metrics"]:
                    plot_series(ax, data[run["id"]][tag], SERIES_COLORS[i], label=run["title"])
            fig.legend(loc="outside lower center", ncol=LEGEND_COLUMNS, frameon=False, fontsize="small")
        else:
            run = job["runs"][0]
            plot_series(ax, data[run["id"]][tag], SERIES_COLORS[0])
            ax.set_title(textwrap.fill(run["title"], TITLE_WRAP))

        save_figure(fig, (image_path(job, tag), flat_image_path(job, tag)))

# =============================================================================
# MAIN
# =============================================================================

def main():
    runs = read_overview(EXCEL_PATH)
    jobs = build_jobs(runs)
    print(f"{len(runs)} runs in sheet -> {len(jobs)} output folders. Checking logs in '{LOG_ROOT}' ...")

    data, issues = check_all(runs, jobs)
    if issues:
        print(f"\n{len(issues)} issue(s) found. No images saved.\n")
        for issue in issues:
            print(f"  - {issue}")
        sys.exit(1)

    print("All logs and metrics found. Exporting ...\n")
    for job in jobs:
        export_job(job, data)
    print(f"\nDone: {sum(len(j['metrics']) for j in jobs)} images in '{OUTPUT_ROOT}' (+ copies in '{FLAT_FOLDER}').")


if __name__ == "__main__":
    main()